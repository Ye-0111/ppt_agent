"""PPT 模板接口 - 上传成品模板 / 列表 / 详情 / 删除 / 选模板 AI 创作（SSE）。"""

import json
import logging
import re
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from databases import Database
from pptx import Presentation
from langchain_core.messages import HumanMessage, SystemMessage

from app.core.config import settings
from app.core.database import get_db
from app.core.deps import get_current_user
from app.core.exceptions import BusinessException, ErrorCode, throw_if
from app.core.observability import apply_trace_context, build_trace_metadata
from app.core.sse import sse_event, SSE_HEADERS
from app.schemas.common import BaseResponse, DeleteRequest
from app.schemas.user import LoginUserVO
from app.schemas.events import (
    TASK_CLASSIFIED, AGENT_THINKING, TEXT_DELTA,
    TOOL_CALL, TOOL_RESULT, DONE, ERROR,
)
from app.schemas.ppt_template import TemplateListRequest, TemplateCreateRequest
from app.services.ppt_service import PptService
from app.services.ppt_template_service import PptTemplateService
from app.agents.common.llm import build_llm
from app.tools.ppt_template import render_ppt_template, scan_template_placeholders

router = APIRouter(prefix="/ppt/template", tags=["PPT模板"])
logger = logging.getLogger(__name__)


# --- 辅助 ---

def _safe_stem(stem: str) -> str:
    """把原始文件名 stem 清理成可安全拼接的片段（保留字母数字/CJK/._-）。"""
    s = re.sub(r"[^\w一-龥.-]", "_", stem).strip("_")
    return s[:40]


def _parse_json_object(text: str):
    """从容错文本里解析 JSON 对象；失败返回 None。"""
    text = text.strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    # ```json ... ``` 围栏
    fence = text.find("```")
    if fence != -1:
        inner_start = text.find("\n", fence)
        inner_end = text.find("```", inner_start + 1 if inner_start != -1 else fence + 3)
        if inner_start != -1 and inner_end > inner_start:
            candidate = text[inner_start + 1:inner_end].strip()
            try:
                obj = json.loads(candidate)
                if isinstance(obj, dict):
                    return obj
            except json.JSONDecodeError:
                pass
    # 首个 { 到末个 }
    s, e = text.find("{"), text.rfind("}")
    if s != -1 and e > s:
        try:
            obj = json.loads(text[s:e + 1])
            if isinstance(obj, dict):
                return obj
        except json.JSONDecodeError:
            pass
    return None


async def _stream_fill_values(
    message: str, placeholders: list[str], trace_meta: dict | None = None,
):
    """LLM 按主题流式生成 {占位名: 值}，逐 chunk yield 文本。"""
    llm = build_llm(temperature=0.6)
    names = ", ".join(placeholders)
    prompt = (
        f"用户想用 PPT 模板生成一份演示文稿，主题/需求如下：\n{message}\n\n"
        f"模板里有以下占位符（写作 {{name}}），请根据主题为每个占位符生成简短贴切的内容"
        f"（每个值不超过 30 字，直接给最终展示文本，不要解释、不要引号）。\n"
        f"占位符列表：{names}\n\n"
        f"只输出一个 JSON 对象，键为占位符名、值为生成内容，不要输出 JSON 以外的任何文字。"
    )
    msgs = [
        SystemMessage(content="你是 PPT 内容填写助手，严格只输出一个 JSON 对象。"),
        HumanMessage(content=prompt),
    ]
    # 纯 LLM 调用无继承的 graph contextvar，显式 config 是挂业务 metadata 的唯一手段
    config = apply_trace_context(
        None,
        metadata={"intent": "template", **(trace_meta or {})},
        tags=["ppt-create", "template"],
        run_name="模板填充",
    )
    async for chunk in llm.astream(msgs, config=config):
        text = chunk.content
        if isinstance(text, str) and text:
            yield text


async def _mark_terminal(service: PptService, ppt_id: int, state: dict, tag: str):
    try:
        if state["filename"]:
            await service.mark_success(ppt_id)
        else:
            await service.mark_failed(ppt_id, "未生成文件")
    except Exception as e:
        logger.error("[%s] 标记终态失败: %s", tag, e)


async def _mark_failed_safe(service: PptService, ppt_id: int, err: str):
    try:
        await service.mark_failed(ppt_id, err)
    except Exception:
        pass


# --- 接口 ---

@router.post("/upload", response_model=BaseResponse)
async def upload_template(
    file: UploadFile = File(..., description="含 {{占位符}} 的 .pptx 模板"),
    name: str = Form(..., description="模板名称"),
    is_public: bool = Form(False, description="是否公开"),
    user: LoginUserVO = Depends(get_current_user),
    db: Database = Depends(get_db),
):
    """上传成品模板（仅接受已含 {{占位符}} 的 PPTX）。

    服务端生成存储文件名（无穿越风险），扫描占位符，入库。
    无占位符 / 解析失败 → 拒绝并清理已落盘文件。
    """
    orig = file.filename or ""
    throw_if(not orig.lower().endswith(".pptx"), ErrorCode.PARAMS_ERROR, "仅支持 .pptx 文件")
    data = await file.read()
    throw_if(not data, ErrorCode.PARAMS_ERROR, "上传文件为空")

    templates_dir = settings.workspace_path / "ppt_templates"
    templates_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"tpl_{user.id}_{uuid.uuid4().hex[:8]}_{_safe_stem(Path(orig).stem)}.pptx"
    stored_path = templates_dir / stored_name
    stored_path.write_bytes(data)

    try:
        prs = Presentation(str(stored_path))
    except Exception:
        stored_path.unlink(missing_ok=True)
        raise BusinessException(ErrorCode.PARAMS_ERROR, "无法解析的 PPTX 文件")

    placeholders = scan_template_placeholders(stored_path)
    if not placeholders:
        stored_path.unlink(missing_ok=True)
        raise BusinessException(
            ErrorCode.PARAMS_ERROR, "模板未检测到 {{占位符}}，请上传含 {{占位符}} 的成品模板"
        )

    service = PptTemplateService(db)
    tpl_id = await service.create(
        user_id=user.id,
        name=name,
        filename=stored_name,
        variables_json=json.dumps(placeholders, ensure_ascii=False),
        slide_count=len(prs.slides),
        is_public=1 if is_public else 0,
    )
    tpl = await service.get(tpl_id, user.id)
    return BaseResponse.success(data=tpl, message="模板上传成功")


@router.post("/list", response_model=BaseResponse)
async def list_templates(
    req: TemplateListRequest,
    user: LoginUserVO = Depends(get_current_user),
    db: Database = Depends(get_db),
):
    """模板列表（分页；scope: mine/public/all，默认 all=自己+公开）。"""
    service = PptTemplateService(db)
    result = await service.list_templates(user.id, req.current, req.page_size, req.scope or "all")
    return BaseResponse.success(data=result)


@router.get("/detail/{template_id}", response_model=BaseResponse)
async def template_detail(
    template_id: int,
    user: LoginUserVO = Depends(get_current_user),
    db: Database = Depends(get_db),
):
    """模板详情（含占位符列表；仅自己或公开可见）。"""
    service = PptTemplateService(db)
    tpl = await service.get(template_id, user.id)
    if not tpl:
        raise BusinessException(ErrorCode.NOT_FOUND_ERROR, "模板不存在或无权访问")
    return BaseResponse.success(data=tpl)


@router.post("/delete", response_model=BaseResponse)
async def delete_template(
    req: DeleteRequest,
    user: LoginUserVO = Depends(get_current_user),
    db: Database = Depends(get_db),
):
    """删除模板（仅 owner；同时清理 ppt_templates/ 下文件）。"""
    service = PptTemplateService(db)
    ok = await service.delete(req.id, user.id)
    if not ok:
        raise BusinessException(ErrorCode.NOT_FOUND_ERROR, "模板不存在或无权删除")
    return BaseResponse.success(message="模板已删除")


@router.post("/create")
async def create_from_template(
    req: TemplateCreateRequest,
    user: LoginUserVO = Depends(get_current_user),
    db: Database = Depends(get_db),
):
    """选模板创作 PPT（SSE）：AI 按主题自动填占位符 → 渲染 → 扣配额 → 落 sys_ppt。

    与 /ppt/create 体验一致：配额门控在开流前；生成成功扣 1 次；DONE 带 pptId/filename/remaining_quota。
    """
    ppt_service = PptService(db)
    tpl_service = PptTemplateService(db)

    # 1. 同步前置（开流前抛业务异常）
    await ppt_service.check_quota(user.id)
    tpl = await tpl_service.get_owned(req.template_id, user.id)
    if not tpl:
        raise BusinessException(ErrorCode.NOT_FOUND_ERROR, "模板不存在或无权使用")
    placeholders = tpl["variables"] or []
    if not placeholders:
        raise BusinessException(ErrorCode.PARAMS_ERROR, "模板无占位符，无法创作")

    topic = req.message[:500]
    style = req.style or "business"
    session_id = f"u{user.id}"
    ppt_id = await ppt_service.create_task(user.id, topic, style, session_id)
    state = {"remaining_quota": None, "deducted": False, "filename": None}
    trace_meta = build_trace_metadata(
        session_id=session_id, user_id=user.id, ppt_id=ppt_id, style=style,
    )

    # 2. SSE generator
    async def event_generator():
        try:
            yield sse_event(TASK_CLASSIFIED, {
                "mode": "template",
                "templateId": tpl["id"],
                "templateName": tpl["name"],
                "slideCount": tpl["slideCount"],
                "placeholders": placeholders,
            })
            yield sse_event(AGENT_THINKING, {
                "role": "content",
                "message": "正在按主题为模板占位符生成内容",
            })

            # AI 流式填值
            parts: list[str] = []
            async for delta in _stream_fill_values(
                req.message, placeholders, trace_meta=trace_meta,
            ):
                parts.append(delta)
                yield sse_event(TEXT_DELTA, {"content": delta})

            values = _parse_json_object("".join(parts))
            if not isinstance(values, dict):
                raise RuntimeError("AI 填充结果解析失败")

            # 渲染（模板从 ppt_templates/ 读，产物写 ppt_output/）
            rendered_name = f"tpl_render_{user.id}_{uuid.uuid4().hex[:8]}.pptx"
            yield sse_event(TOOL_CALL, {"tool": "render_ppt_template"})
            r = await render_ppt_template.ainvoke({
                "template_filename": tpl["filename"],
                "variables": json.dumps(values, ensure_ascii=False),
                "output_filename": rendered_name,
                "source_dir": "ppt_templates",
            })
            render_res = json.loads(r)
            if not render_res.get("success"):
                raise RuntimeError(render_res.get("message", "模板渲染失败"))

            # 扣配额
            _ok, remaining = await ppt_service.deduct_for_task(ppt_id, user.id, rendered_name)
            state["deducted"] = True
            state["remaining_quota"] = remaining
            state["filename"] = rendered_name

            yield sse_event(TOOL_RESULT, {
                "tool": "render_ppt_template",
                "substituted": render_res.get("substituted", []),
                "missing": render_res.get("missing_variables", []),
            })

            done_data = {"pptId": ppt_id, "filename": rendered_name}
            if remaining is not None:
                done_data["remaining_quota"] = remaining
            yield sse_event(DONE, done_data)

            await _mark_terminal(ppt_service, ppt_id, state, "模板创作")
        except Exception as e:
            logger.error("[模板创作] 未捕获异常: %s", e)
            await _mark_failed_safe(ppt_service, ppt_id, str(e))
            yield sse_event(ERROR, {"message": f"服务器异常: {e}"})

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers=SSE_HEADERS,
    )
